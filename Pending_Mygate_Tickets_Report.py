# Pending MyGate Tickets Report
#
# Replace with your own MyGate API token.
#
import requests
import time
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

TOKEN = "zr2Er9wrdfhTyiY01Lnvr03cje9oeeH7wR6XvFmeSR87okw1qW4QyAuRkoSaOkff"
URL = "https://api.dashboard.mygate.com/graphql/"

HEADERS = {
    "authorization": TOKEN,
    "content-type": "application/json",
    "origin": "https://dashboard.mygate.com",
    "referer": "https://dashboard.mygate.com/"
}

REQUEST_DELAY = 1

CATEGORIES = [
    ("Accounts Billing",252434),
    ("Construction Or Project Related",277747),
    ("Design Related Issue",277632),
    ("Estate Infra Outer Area from the plot",267181),
    ("FM Common Area Related Issue",277745),
    ("IT WIFI Network",277816),
    ("Products Appliances",277744),
]

TOTAL=["open","hold","re_opened","job_done","in_progress","closed"]
OPEN=["open","hold","re_opened","job_done","in_progress"]
RESOLVED=["closed"]

def epoch(s):
    return int(datetime.strptime(s,"%d-%m-%Y").timestamp())

def get_count(cat, from_date, to_date, statuses):
    payload={
        "operationName":"getAdminSrList",
        "variables":{"requestData":{
            "requiredFields":["id"],
            "pagination":{"count":1,"page":1},
            "sorting":[],
            "conditions":[
                {"name":"date_filter","operation":"equal","values":["created_date"]},
                {"name":"category","values":[cat],"operation":"equal"},
                {"name":"from_date","values":[epoch(from_date)],"operation":"gte"},
                {"name":"to_date","values":[epoch(to_date)+86399],"operation":"lte"},
                {"name":"mygate_status","values":statuses,"operation":"equal" if len(statuses)==1 else "in"}
            ]}},
        "query":"query getAdminSrList($requestData: DataListInput){getAdminSrList(requestData:$requestData){dataResponse{totalCount}}}"
    }

    r=requests.post(URL,headers=HEADERS,json=payload,timeout=60)
    r.raise_for_status()
    data=r.json()
    time.sleep(REQUEST_DELAY)
    return data["data"]["getAdminSrList"]["dataResponse"]["totalCount"]

from_date="01-01-2024"

today=datetime.now().strftime("%d-%m-%Y")
to_date=input(f"Report till date (DD-MM-YYYY) [Press Enter for Today ({today})]: ").strip()
if not to_date:
    to_date=today

wb=Workbook()
ws=wb.active
ws.title="Pending Tickets"

BLUE="4D93D9"
blue_fill=PatternFill(fill_type="solid",start_color=BLUE,end_color=BLUE)

title_font=Font(name="Aptos",size=12,bold=False,color="000000")
header_font=Font(name="Aptos",size=12,bold=False,color="000000")
body_font=Font(name="Aptos",size=12,color="000000")

thin=Side(border_style="thin",color="000000")
border=Border(left=thin,right=thin,top=thin,bottom=thin)

ws.merge_cells("A1:D1")
title=ws["A1"]
title.value=f"Pending Mygate Tickets - From {from_date.lstrip('0').replace('-0','-')} To {to_date.lstrip('0').replace('-0','-')}"
title.fill=blue_fill
title.font=title_font
title.alignment=Alignment(horizontal="center",vertical="center")

headers=["Category","Total","Resolved","Open"]
for col,h in enumerate(headers,1):
    c=ws.cell(row=2,column=col)
    c.value=h
    c.fill=blue_fill
    c.font=header_font
    c.alignment=Alignment(horizontal="left",vertical="center")

row=3
tt=tr=to=0

for name,cid in CATEGORIES:
    print(f"Fetching: {name}")
    total=get_count(cid,from_date,to_date,TOTAL)
    resolved=get_count(cid,from_date,to_date,RESOLVED)
    openv=get_count(cid,from_date,to_date,OPEN)

    vals=[name,total,resolved,openv]
    for col,val in enumerate(vals,1):
        cell=ws.cell(row=row,column=col)
        cell.value=val
        cell.font=body_font
        cell.border=border
        if col==1:
            cell.alignment=Alignment(horizontal="left",vertical="center")
        else:
            cell.alignment=Alignment(horizontal="right",vertical="center")

    tt+=total
    tr+=resolved
    to+=openv
    row+=1

totals=["Total",tt,tr,to]
for col,val in enumerate(totals,1):
    cell=ws.cell(row=row,column=col)
    cell.value=val
    cell.font=body_font
    cell.border=border
    if col==1:
        cell.alignment=Alignment(horizontal="left",vertical="center")
    else:
        cell.alignment=Alignment(horizontal="right",vertical="center")

for r in ws.iter_rows(min_row=1,max_row=row,max_col=4):
    for cell in r:
        cell.border=border

for column_cells in ws.iter_cols(min_col=1,max_col=4):
    letter=get_column_letter(column_cells[0].column)
    max_len=max(len(str(c.value or "")) for c in column_cells)
    ws.column_dimensions[letter].width=max_len+4

filename=f"Pending_Mygate_Tickets_Report_{to_date}.xlsx"
wb.save(filename)
print(f"Report saved as {filename}")
